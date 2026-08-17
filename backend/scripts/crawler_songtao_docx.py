# -*- coding: utf-8 -*-
"""松桃县事故调查报告 docx 批量爬虫：页面列出的调查报告 docx 附件 → python-docx 提取。

列表：https://www.songtao.gov.cn/zwgk/xxgkml/jcxxgk/yjgl/202605/t20260521_90195273.html
（2023-2025 年事故调查报告公开页，附件为 .docx）
"""
import io
import json
import re
import sys
import time
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(Path(__file__).resolve().parent))

import requests

from crawler_safehoo import OUTPUT_DIR, fetch, save_doc  # noqa: E402

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0 Safari/537.36"}
LIST_URL = "https://www.songtao.gov.cn/zwgk/xxgkml/jcxxgk/yjgl/202605/t20260521_90195273.html"


def docx_to_paras(data: bytes) -> list[str]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paras: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if len(t) >= 12:
            paras.append(t)
    # 表格内容也提取（报告含大量表格）
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                line = " | ".join(cells)
                if len(line) >= 12:
                    paras.append(line)
    return paras


def main() -> int:
    html = fetch(LIST_URL)
    if not html:
        print("列表页抓取失败")
        return 1
    # 附件链接：.docx
    atts = re.findall(r'href="([^"]+\.docx)"[^>]*>([^<]{4,80})</a>', html)
    print(f"发现 {len(atts)} 个 docx 附件")
    total = 0
    for href, label in atts:
        title = label.strip()
        if not title or "事故" not in title:
            continue
        url = href if href.startswith("http") else (
            "https://www.songtao.gov.cn" + href if href.startswith("/") else
            LIST_URL.rsplit("/", 1)[0] + "/" + href.lstrip("./"))
        safe = re.sub(r'[\\/:*?"<>|]', "_", title)[:80]
        if (OUTPUT_DIR / f"{safe}.json").exists():
            continue
        try:
            r = requests.get(url, headers=UA, timeout=60)
            if r.status_code != 200 or len(r.content) < 500:
                print(f"  ✗ 下载失败 {title[:30]}")
                continue
            paras = docx_to_paras(r.content)
        except Exception as e:
            print(f"  ✗ 解析失败 {title[:30]}: {e}")
            continue
        if len(paras) < 10:
            print(f"  ✗ 内容不足 {title[:30]}（{len(paras)} 段）")
            continue
        chapters = [{"chapter": "正文", "articles": [{"no": f"第{i+1}条", "content": p}
                                                     for i, p in enumerate(paras)]}]
        if save_doc(title, "case", "典型事故案例", url, chapters, source="政府网站公开"):
            total += 1
            print(f"  ✓ {title[:40]}（{len(paras)} 条）")
        time.sleep(0.5)
    print(f"完成：新增 {total} 篇")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
